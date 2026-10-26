"""Evidence layer: the source registry and the harmonised published-curve library.

Each source carries a licence class, because no commercial product may assume a
published curve is redistributable. Classes follow the product specification:

    OPEN              free to redistribute and to train on
    REFERENCE_ONLY    may be cited and benchmarked against, not redistributed
    LICENSED          requires a commercial agreement before ingestion
    CLIENT_PROPRIETARY supplied by a client, never leaves the tenant
    RESTRICTED        known to exist, not usable under current terms
"""
from __future__ import annotations

import hashlib
import os
from dataclasses import dataclass, asdict

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
JRC_XLSX = os.path.join(HERE, "..", "data", "raw", "jrc_global_ddf.xlsx")
CURVES_CSV = os.path.join(HERE, "data", "jrc_curves.csv")
MAXDAMAGE_CSV = os.path.join(HERE, "data", "jrc_max_damage.csv")


@dataclass(frozen=True)
class Source:
    source_id: str
    title: str
    authors: str
    year: int
    geography: str
    approach: str
    licence_class: str
    citation: str
    obtained: str          # how this repository actually got it
    status: str            # INGESTED | BENCHMARK_ONLY | NOT_OBTAINED


REGISTRY: tuple[Source, ...] = (
    Source(
        "JRC_GLOBAL_2017",
        "Global flood depth-damage functions: methodology and the database with guidelines",
        "Huizinga, J.; de Moel, H.; Szewczyk, W.", 2017,
        "Global, 6 continents, 214 countries", "Normalised continental curves + country max damage",
        "OPEN",
        "JRC Technical Report EUR 28552 EN. doi:10.2760/16510",
        "Downloaded from the JRC publications repository as the published .xlsx database",
        "INGESTED",
    ),
    Source(
        "NFIP_CLAIMS_V3",
        "NFIP Redacted Claims v3 (OpenFEMA)",
        "Federal Emergency Management Agency", 2026,
        "United States", "Empirical, object level, adjuster assessed",
        "OPEN",
        "FEMA OpenFEMA dataset NfipClaims v3, data frozen 2026-06-01",
        "Downloaded via the OpenFEMA API with keyset pagination",
        "INGESTED",
    ),
    Source(
        "HAZUS_FLOOD",
        "Hazus Flood Model Technical Manual (depth-damage function library)",
        "Federal Emergency Management Agency", 2025,
        "United States", "Empirical/synthetic by occupancy, storeys, foundation",
        "OPEN",
        "FEMA Hazus 7.0 Flood Model Technical Manual, June 2025",
        "Document endpoint returned HTTP 403 to automated retrieval in this environment",
        "NOT_OBTAINED",
    ),
    Source(
        "USACE_NACCS",
        "North Atlantic Coast Comprehensive Study physical depth-damage functions",
        "US Army Corps of Engineers", 2015,
        "US North Atlantic coast", "Empirical, triangular min/most-likely/max",
        "OPEN",
        "USACE NACCS, Physical Depth Damage Function Summary Report, 26 Jan 2015",
        "Document endpoint returned HTTP 403 to automated retrieval in this environment",
        "NOT_OBTAINED",
    ),
    Source(
        "HOWAS21",
        "HOWAS21 object-specific flood damage database",
        "Kellermann, P.; Schroeter, K.; Thieken, A. H.; Haubrock, S.-N.; Kreibich, H.", 2020,
        "Germany, Austria", "Object-level empirical database",
        "REFERENCE_ONLY",
        "Nat. Hazards Earth Syst. Sci. 20, 2503-2519. doi:10.5194/nhess-20-2503-2020",
        "Access requires registration and a data agreement; not automatable here",
        "NOT_OBTAINED",
    ),
    Source(
        "MCM_UK",
        "Multi-Coloured Manual property and economic loss functions",
        "Penning-Rowsell, E. et al.", 2013,
        "United Kingdom", "Detailed synthetic property loss functions",
        "LICENSED",
        "Flood Hazard Research Centre, Middlesex University",
        "Commercial licence required before ingestion; deliberately not scraped",
        "NOT_OBTAINED",
    ),
    Source(
        "INSYDE",
        "INSYDE: a synthetic, probabilistic flood damage model based on explicit cost analysis",
        "Dottori, F.; Figueiredo, R.; Martina, M. L. V.; Molinari, D.; Scorzini, A. R.", 2016,
        "Italy, adaptable", "Component-based synthetic engineering model",
        "OPEN",
        "Nat. Hazards Earth Syst. Sci. 16, 2577-2591. doi:10.5194/nhess-16-2577-2016",
        "Method reimplementable from the paper; not ingested in this build",
        "BENCHMARK_ONLY",
    ),
)


def registry_frame() -> pd.DataFrame:
    return pd.DataFrame([asdict(s) for s in REGISTRY])


# --------------------------------------------------------------------------
# JRC ingestion
# --------------------------------------------------------------------------
JRC_CONTINENTS = ["EUROPE", "North AMERICA", "Centr&South AMERICA",
                  "ASIA", "AFRICA", "OCEANIA", "GLOBAL"]


def ingest_jrc(xlsx_path: str = JRC_XLSX) -> pd.DataFrame:
    """Read the published JRC workbook into a tidy long curve table.

    Returns columns: damage_class, continent, depth_m, damage_factor.
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb["Damage functions"]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=10, values_only=True))

    records = []
    current = None
    for r in rows:
        label = (str(r[0]).strip() if r[0] is not None else "")
        if label and not label.lower().startswith(("damage class", "flood depth")):
            current = label
        depth = r[1]
        if current is None or not isinstance(depth, (int, float)):
            continue
        for j, cont in enumerate(JRC_CONTINENTS):
            v = r[2 + j] if 2 + j < len(r) else None
            if isinstance(v, (int, float)):
                records.append((current, cont, float(depth), float(v)))

    df = pd.DataFrame(records, columns=["damage_class", "continent", "depth_m", "damage_factor"])
    df = df[df["damage_class"].str.len() > 2].reset_index(drop=True)
    return df


def ingest_jrc_max_damage(xlsx_path: str = JRC_XLSX) -> pd.DataFrame:
    """Country-level maximum damage values for the residential class (EUR/m2)."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb["MaxDamage-Residential"]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=6, values_only=True))
    header_idx = None
    for i, r in enumerate(rows):
        joined = " ".join(str(c) for c in r if c is not None).lower()
        if "country" in joined:
            header_idx = i
            break
    if header_idx is None:
        return pd.DataFrame()
    header = [str(c).strip() if c is not None else f"c{j}" for j, c in enumerate(rows[header_idx])]
    body = [r for r in rows[header_idx + 1:] if r[0] is not None]
    return pd.DataFrame(body, columns=header)


def curve_interpolator(curves: pd.DataFrame, damage_class: str, continent: str):
    """Return f(depth_m) -> damage factor for one published curve.

    Linear between the published knots, flat above the deepest knot. The JRC
    curves are defined at 0, 0.5, 1, 1.5, 2, 3, 4, 5, 6 m above floor level.
    """
    sel = curves[(curves.damage_class == damage_class) & (curves.continent == continent)]
    sel = sel.sort_values("depth_m")
    if sel.empty:
        raise KeyError(f"no JRC curve for {damage_class} / {continent}")
    x = sel["depth_m"].to_numpy()
    y = sel["damage_factor"].to_numpy()

    def f(d):
        d = np.asarray(d, dtype=float)
        return np.interp(np.clip(d, x[0], x[-1]), x, y)

    return f


def provenance_hash(*paths: str) -> str:
    """Content hash over the evidence snapshot, pinned into every prediction."""
    h = hashlib.sha256()
    for p in sorted(paths):
        if os.path.exists(p):
            h.update(os.path.basename(p).encode())
            with open(p, "rb") as fh:
                for chunk in iter(lambda: fh.read(1 << 20), b""):
                    h.update(chunk)
    return h.hexdigest()[:16]


def build_curve_library(out_csv: str = CURVES_CSV) -> pd.DataFrame:
    os.makedirs(os.path.dirname(out_csv), exist_ok=True)
    df = ingest_jrc()
    df.to_csv(out_csv, index=False)
    return df


def load_curves(path: str = CURVES_CSV) -> pd.DataFrame:
    if not os.path.exists(path):
        return build_curve_library(path)
    return pd.read_csv(path)
